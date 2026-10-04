"""CSV / Excel / PDF exports of a research job (professor contact list)."""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models as m
from app.services.results import professors_query, sources_for, university_payload

COLUMNS = ["Professor", "Chinese Name", "Position", "Department", "Email", "Email Status", "Profile URL",
           "Source URL", "Verification Status"]


def _prof_source(db: Session, p: m.Professor) -> str:
    srcs = sources_for(db, p.job_id, "professor", p.id)
    return srcs[0].source_url if srcs else ""


def _row(db: Session, p: m.Professor) -> list:
    return [p.name, p.name_chinese or "", p.position or "", p.department_name or "",
            p.email or "Not publicly listed", p.email_status, p.profile_url or "", _prof_source(db, p),
            p.verification_status]


def export_csv(db: Session, job_id: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    for p in db.scalars(professors_query(job_id)).all():
        w.writerow(_row(db, p))
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens Chinese correctly


def export_excel(db: Session, job_id: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    job = db.get(m.ResearchJob, job_id)
    uni = university_payload(db, job)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(title: str, headers: list[str], rows: list[list], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = header_font, header_fill
        for r in rows:
            ws.append(["" if v is None else v for v in r])
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [min(len(str(r[i - 1] or "")), 60) for r in rows]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 62)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    profs = db.scalars(professors_query(job_id)).all()
    sheet("Professors", COLUMNS, [_row(db, p) for p in profs], first=True)
    sheet("University", ["Attribute", "Value"], [
        ["Name", uni.name if uni else "Not Found"], ["Chinese name", uni.name_chinese if uni else ""],
        ["Official URL", uni.official_url if uni else job.university_url], ["Country", uni.country if uni else ""],
        ["Location", uni.location if uni else ""], ["Research fields", ", ".join(job.fields or [])],
        ["Professors found", len(profs)], ["With public email", sum(1 for p in profs if p.email)],
        ["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")],
    ])
    depts = db.scalars(select(m.Department).where(m.Department.job_id == job_id)).all()
    sheet("Departments", ["Department", "Chinese name", "Website", "Faculty list URLs", "Matched fields"],
          [[d.name, d.name_chinese, d.url, "\n".join(d.faculty_list_urls or []), ", ".join(d.matched_fields or [])]
           for d in depts])
    srcs = db.scalars(select(m.Source).where(m.Source.job_id == job_id).order_by(m.Source.id)).all()
    sheet("Sources", ["Entity", "Entity ID", "Supports", "URL", "Type", "Title", "Retrieved at"],
          [[x.entity_type, x.entity_id, x.supports, x.source_url, x.source_type, x.title,
            x.retrieved_at.strftime("%Y-%m-%d %H:%M") if x.retrieved_at else ""] for x in srcs])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_pdf(db: Session, job_id: str) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    try:  # built-in CID font renders Chinese without shipping font files
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        cjk = "STSong-Light"
    except Exception:  # pragma: no cover
        cjk = "Helvetica"

    styles = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=styles["BodyText"], fontName=cjk, fontSize=8.5, leading=11)
    small = ParagraphStyle("s", parent=body, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#555555"))
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontName=cjk, textColor=colors.HexColor("#1F3A5F"))

    def P(text, st=body):
        return Paragraph(escape(str(text or "")).replace("\n", "<br/>"), st)

    job = db.get(m.ResearchJob, job_id)
    uni = university_payload(db, job)
    profs = db.scalars(professors_query(job_id)).all()
    title = (uni.name if uni else job.university_url) or job.university_url
    if uni and uni.name_chinese and uni.name_chinese != uni.name:
        title += f" ({uni.name_chinese})"
    story = [P(f"Professor contacts — {title}", h1),
             P(f"Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC · Fields: {', '.join(job.fields or [])} · "
               f"{len(profs)} professors, {sum(1 for p in profs if p.email)} with a public email", small),
             Spacer(1, 6)]
    data = [[P("Professor"), P("Position"), P("Department"), P("Email"), P("Profile"), P("Status")]]
    for p in profs:
        name = p.name + (f"\n{p.name_chinese}" if p.name_chinese and p.name_chinese != p.name else "")
        data.append([P(name), P(p.position or "—"), P(p.department_name or "—"),
                     P(p.email or "Not publicly listed"), P(p.profile_url or "—", small), P(p.verification_status)])
    story.append(Table(data, colWidths=[40 * mm, 32 * mm, 50 * mm, 55 * mm, 75 * mm, 25 * mm], repeatRows=1,
                       style=TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BBBBBB")),
                                         ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F7")),
                                         ("VALIGN", (0, 0), (-1, -1), "TOP")])))
    if job.warnings:
        story += [Spacer(1, 8), P("Notes & limitations", body)] + [P("• " + w, small) for w in job.warnings]
    story += [Spacer(1, 8), P("Emails and names come only from the cited official pages. Always confirm contact "
                              "details with the university before writing.", small)]
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                      bottomMargin=12 * mm, title="Professor contacts").build(story)
    return out.getvalue()
